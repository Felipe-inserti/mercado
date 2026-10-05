"""Imagens do README (Sprint 23) -- gráficos vivem em notebook, não em `src/`
(CLAUDE.md, seção 2). Lógica de negócio nenhuma: só leitura dos resultados já em
disco (as mesmas funções do relatório mensal) e desenho.

    python notebooks/gerar_imagens_readme.py

Gera em `docs/img/`:

1. `servico_e_capital.png` -- dois gráficos pequenos empilhados (serviço em % e capital
   em índice, 100 = baseline recalibrado). Nunca um eixo duplo: duas medidas de escalas
   diferentes são dois gráficos.
2. `lista_compra_exemplo.png` -- a lista principal de um fornecedor, REDESENHADA a partir do
   xlsx gerado (`results/lista_compra/2017-03-06/`). Não é captura de tela, e a coluna de
   valor em R$ fica de fora (R$ só como ilustrativo).
3. `relatorio_mensal_resumo.png` -- a aba "Resumo visual" de `docs/entregaveis/relatorio_
   mensal.xlsx`, REDESENHADA a partir das células (também não é captura de tela).

Cores: três primeiros slots da paleta categórica de referência, validados com
`validate_palette.js --pairs all` (CVD ΔE 9,2; visão normal 24,0). O verde-água fica abaixo de
3:1 de contraste na superfície clara, então todo valor aparece rotulado na própria barra.
Uma cor por ENTIDADE, a mesma nos dois gráficos.
"""

from __future__ import annotations

from pathlib import Path
from typing import Final

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook

from motor.experiments.iso_service import CELL_LEAD_TIME_DAYS, CELL_REVIEW_PERIOD_DAYS
from motor.reporting.monthly_report import (
    STALE_LABEL_HEADLINE,
    load_motor_financials,
    load_recalibrado_financials,
    load_stale_financials,
)

OUT_DIR: Final[Path] = Path("docs/img")
LISTA_XLSX: Final[Path] = Path("results/lista_compra/2017-03-06/SUP-CLEANING.xlsx")
RELATORIO_XLSX: Final[Path] = Path("docs/entregaveis/relatorio_mensal.xlsx")

SURFACE: Final[str] = "#fcfcfb"
INK: Final[str] = "#0b0b0b"
INK_2: Final[str] = "#52514e"
ORANGE: Final[str] = "#eb6834"  # ERP desatualizado
BLUE: Final[str] = "#2a78d6"  # ERP recalibrado
AQUA: Final[str] = "#1baf7a"  # motor
CATEGORIES: Final[tuple[str, ...]] = ("ERP\ndesatualizado", "ERP\nrecalibrado", "Motor\n(braço 2)")
COLORS: Final[tuple[str, ...]] = (ORANGE, BLUE, AQUA)


def _dados() -> dict[str, tuple[float, float, float]]:
    stale = load_stale_financials(STALE_LABEL_HEADLINE)
    recal = load_recalibrado_financials()
    motor = load_motor_financials()
    arms = (stale, recal, motor)
    return {
        "servico": tuple(a.nivel_servico * 100 for a in arms),  # type: ignore[dict-item]
        "capital": tuple(100.0 * a.capital_medio_rs / recal.capital_medio_rs for a in arms),  # type: ignore[dict-item]
    }


def _bars(ax: plt.Axes, values: tuple[float, ...], *, title: str, fmt: str, ymax: float) -> None:
    ax.set_facecolor(SURFACE)
    bars = ax.bar(CATEGORIES, values, color=COLORS, width=0.5)
    for bar, value in zip(bars, values, strict=True):
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            value + ymax * 0.02,
            fmt.format(value),
            ha="center",
            va="bottom",
            fontsize=10,
            color=INK,
        )
    ax.set_ylim(0, ymax)
    ax.set_title(title, fontsize=10, color=INK, loc="left")
    ax.tick_params(axis="x", labelsize=8.5, colors=INK_2, length=0)
    ax.tick_params(axis="y", labelsize=8, colors=INK_2, length=0)
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    ax.spines["bottom"].set_color("#c9c8c2")
    ax.set_yticks([])


def _servico_e_capital(path: Path) -> None:
    d = _dados()
    fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(5.2, 6.2), facecolor=SURFACE)
    _bars(ax1, d["servico"], title="Nível de serviço (%)", fmt="{:.1f}%", ymax=110)
    _bars(
        ax2,
        d["capital"],
        title="Capital médio empregado (índice, 100 = baseline recalibrado)",
        fmt="{:.0f}",
        ymax=125,
    )
    fig.text(
        0.02,
        0.015,
        f"Célula lead time {CELL_LEAD_TIME_DAYS} / revisão {CELL_REVIEW_PERIOD_DAYS}; "
        "backtest sobre o dado do Favorita.\n"
        "Fonte: results/baseline_desatualizado, results/sensibilidade_dirigida/lt7_rp14,\n"
        "results/iso_servico (braço 2, alpha 0,66).",
        fontsize=7.5,
        color=INK_2,
        va="bottom",
    )
    fig.tight_layout(rect=(0, 0.09, 1, 1))
    fig.savefig(path, dpi=160, facecolor=SURFACE)
    plt.close(fig)


def _lista_compra(path: Path) -> None:
    ws = load_workbook(LISTA_XLSX)["Lista principal"]
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=7)]
    header = rows[0]
    keep = [0, 1, 3, 4, 5, 6, 7]  # tudo menos a unidade e o valor em R$
    names = ["código", "categoria", "sugestão\n(unid.)", "sugestão\n(fardos)", "estoque\natual",
             "em\ntrânsito", "cobertura\n(dias)"]  # fmt: skip
    body = []
    for r in rows[1:]:
        cells = [r[i] for i in keep]
        body.append(
            [
                f"{v:,.0f}"
                if isinstance(v, (int, float)) and i in (2, 3, 4, 5)
                else (f"{v:.1f}" if isinstance(v, float) else str(v))
                for i, v in enumerate(cells)
            ]
        )
    assert header[0] == "código"
    fig, ax = plt.subplots(figsize=(8.4, 3.6), facecolor=SURFACE)
    ax.axis("off")
    ax.set_title(
        "Lista principal -- fornecedor SUP-CLEANING, decisão em 2017-03-06 (dado do Favorita)",
        fontsize=10,
        color=INK,
        loc="left",
        pad=10,
    )
    table = ax.table(cellText=body, colLabels=names, cellLoc="center", bbox=(0.0, 0.0, 1.0, 0.88))
    table.auto_set_font_size(False)
    table.set_fontsize(9)
    for (r, _c), cell in table.get_celld().items():
        cell.set_edgecolor("#e3e2dc")
        if r == 0:
            cell.set_facecolor("#1f3864")
            cell.get_text().set_color("white")
    fig.text(
        0.01,
        0.02,
        "Redesenhada a partir do xlsx gerado (results/lista_compra/2017-03-06/SUP-CLEANING.xlsx); "
        "a coluna de valor em R\\$ foi omitida.",
        fontsize=7.5,
        color=INK_2,
    )
    fig.savefig(path, dpi=160, facecolor=SURFACE, bbox_inches="tight")
    plt.close(fig)


def _relatorio(path: Path) -> None:
    ws = load_workbook(RELATORIO_XLSX)["Resumo visual"]
    col = [row[0].value for row in ws.iter_rows()]

    def after(label: str, offset: int = 1) -> object:
        return col[col.index(label) + offset]

    servico = str(after("NÍVEL DE SERVIÇO"))
    capital = float(after("CAPITAL PARADO"))  # type: ignore[arg-type]
    margem = float(after("MARGEM POR REAL INVESTIDO"))  # type: ignore[arg-type]
    d = _dados()

    fig = plt.figure(figsize=(4.2, 9.4), facecolor=SURFACE)
    fig.text(
        0.05, 0.965, "RELATÓRIO MENSAL -- RESUMO", fontsize=12, fontweight="bold", color="#1f3864"
    )
    blocos = (
        ("NÍVEL DE SERVIÇO", servico, INK),
        ("CAPITAL PARADO", f"{capital:+.0%}", INK),
        ("MARGEM POR REAL INVESTIDO", f"{margem:+.1%}".replace(".", ","), INK),
    )
    for i, (label, big, color) in enumerate(blocos):
        y = 0.91 - i * 0.105
        fig.text(0.05, y, label, fontsize=8, fontweight="bold", color=INK_2)
        fig.text(0.05, y - 0.045, big, fontsize=26, fontweight="bold", color=color)
    ax1 = fig.add_axes((0.08, 0.36, 0.84, 0.17))
    ax2 = fig.add_axes((0.08, 0.10, 0.84, 0.17))
    _bars(ax1, d["servico"], title="Nível de serviço (%)", fmt="{:.1f}%", ymax=110)
    _bars(ax2, d["capital"], title="Capital (índice, 100 = recalibrado)", fmt="{:.0f}", ymax=125)
    fig.text(
        0.05,
        0.025,
        'Redesenhado a partir da aba "Resumo visual" de relatorio_mensal.xlsx.\n'
        f"Célula lead time {CELL_LEAD_TIME_DAYS} / revisão {CELL_REVIEW_PERIOD_DAYS}; "
        "só números relativos.",
        fontsize=7,
        color=INK_2,
    )
    fig.savefig(path, dpi=160, facecolor=SURFACE)
    plt.close(fig)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    _servico_e_capital(OUT_DIR / "servico_e_capital.png")
    _lista_compra(OUT_DIR / "lista_compra_exemplo.png")
    _relatorio(OUT_DIR / "relatorio_mensal_resumo.png")
    for name in ("servico_e_capital", "lista_compra_exemplo", "relatorio_mensal_resumo"):
        print(OUT_DIR / f"{name}.png")


if __name__ == "__main__":
    main()

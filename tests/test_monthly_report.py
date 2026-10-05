"""Testes de `motor.reporting.monthly_report` (Sprint 19, Fase 4).

Só a parte pura/de leitura (`n_months_no_periodo`, `mensalizar`,
`ruptura_evitada`, `capital_liberado`, `perda_evitada`,
`sensibilidade_margem`, `load_motor_financials`/`load_recalibrado_financials`/
`load_stale_financials`) -- `itens_sangrando_margem` (lê `config/params.yaml`
e o canônico real) e `write_monthly_report_workbook` (xlsx completo) são
verificados manualmente sobre o resultado real gerado, mesmo espírito de
`tests/test_purchase_list.py` (`generate_purchase_list` fora do escopo
automatizado).

`monthly_report` reusa as funções de carregamento de
`motor.experiments.stale_baseline` (`load_baseline_recalibrado`,
`load_motor_iso_servico`, `load_stale_scenario`) -- essas funções leem os
diretórios GLOBAIS do módulo `stale_baseline` em tempo de chamada, não em
tempo de import, então redirecioná-los (monkeypatch) também redireciona as
chamadas feitas de dentro de `monthly_report`.
"""

from __future__ import annotations

import json
from pathlib import Path

import polars as pl
import pytest
import xlsxwriter
from openpyxl import load_workbook

from motor.experiments import stale_baseline as sb
from motor.reporting import monthly_report as mr


def _write_erp_manifest(
    path: Path,
    *,
    factor: float,
    min_order_units: float,
    nivel_servico: float,
    decision_metric: float,
    capital_medio_rs: float,
    moving_average_weeks: int = 4,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(
            {
                "erp_baseline_calibration": {
                    "factor": factor,
                    "min_order_units": min_order_units,
                    "target_service_level": 0.95,
                    "moving_average_weeks": moving_average_weeks,
                    "coverage_days_implied": factor * moving_average_weeks * 7,
                    "coverage_context": "texto preso em outra célula -- nunca lido aqui",
                },
                "results": {
                    "portfolio": {
                        "nivel_servico": nivel_servico,
                        "decision_metric": decision_metric,
                        "capital_medio_rs": capital_medio_rs,
                    }
                },
                "simulation_window": {
                    "start": "2016-08-14",
                    "warmup_days": 90,
                    "evaluation_start": "2016-11-12",
                    "end": "2017-08-15",
                    "n_evaluation_days": 270,
                    "seed": 42,
                },
            }
        )
    )


def _write_item_metrics(
    path: Path, *, margem_realizada_rs: float, ruptura_rs: float, perda_rs: float = 0.0
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    pl.DataFrame(
        {
            "item_id": ["A", "B"],
            "margem_realizada_rs": [margem_realizada_rs / 2, margem_realizada_rs / 2],
            "ruptura_rs": [ruptura_rs / 2, ruptura_rs / 2],
            "perda_rs": [perda_rs / 2, perda_rs / 2],
        }
    ).write_parquet(path)


@pytest.fixture
def cenario(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    results = tmp_path / "results"
    sensitivity_dir = results / "sensibilidade_dirigida"
    iso_dir = results / "iso_servico"
    stale_dir = results / "baseline_desatualizado"

    # stale_baseline lê estes diretórios (funções reusadas por monthly_report)
    monkeypatch.setattr(sb, "SENSITIVITY_DIR", sensitivity_dir)
    monkeypatch.setattr(sb, "ISO_SERVICE_DIR", iso_dir)
    monkeypatch.setattr(sb, "STALE_BASELINE_DIR", stale_dir)
    # monthly_report tem sua PRÓPRIA cópia dos mesmos caminhos
    monkeypatch.setattr(mr, "SENSITIVITY_DIR", sensitivity_dir)
    monkeypatch.setattr(mr, "ISO_SERVICE_DIR", iso_dir)
    monkeypatch.setattr(mr, "STALE_BASELINE_DIR", stale_dir)

    _write_erp_manifest(
        sensitivity_dir / "lt3_rp7" / "erp_baseline" / "manifest.json",
        factor=0.34,
        min_order_units=0.0,
        nivel_servico=0.60,
        decision_metric=17.2,
        capital_medio_rs=175000.0,
    )
    _write_erp_manifest(
        sensitivity_dir / "lt5_rp7" / "erp_baseline" / "manifest.json",
        factor=0.40,
        min_order_units=3.0,
        nivel_servico=0.9358073,
        decision_metric=17.41777597,
        capital_medio_rs=269_644.36,
    )
    _write_erp_manifest(
        sensitivity_dir / "lt7_rp14" / "erp_baseline" / "manifest.json",
        factor=0.8,
        min_order_units=0.0,
        nivel_servico=0.9485436,
        decision_metric=7.08751197,
        capital_medio_rs=672427.37,
    )
    _write_item_metrics(
        sensitivity_dir / "lt7_rp14" / "erp_baseline" / "item_metrics.parquet",
        margem_realizada_rs=4_765_837.05,
        ruptura_rs=265_066.16,
    )
    _write_erp_manifest(
        stale_dir / "calibrado_lt3_rp7" / "erp_baseline" / "manifest.json",
        factor=0.34,
        min_order_units=0.0,
        nivel_servico=0.5999,
        decision_metric=17.2,
        capital_medio_rs=175_113.51,
    )
    _write_item_metrics(
        stale_dir / "calibrado_lt3_rp7" / "erp_baseline" / "item_metrics.parquet",
        margem_realizada_rs=3_011_854.80,
        ruptura_rs=2_019_048.41,
    )
    _write_erp_manifest(
        stale_dir / "calibrado_lt5_rp7" / "erp_baseline" / "manifest.json",
        factor=0.40,
        min_order_units=3.0,
        nivel_servico=0.6635,
        decision_metric=15.0,
        capital_medio_rs=222_082.20,
    )
    _write_item_metrics(
        stale_dir / "calibrado_lt5_rp7" / "erp_baseline" / "item_metrics.parquet",
        margem_realizada_rs=3_331_827.42,
        ruptura_rs=1_699_075.80,
    )

    iso_dir.mkdir(parents=True, exist_ok=True)
    pl.DataFrame(
        {
            "alpha": [0.66],
            "nivel_servico": [0.9483648],
            "decision_metric": [7.29191612],
            "capital_medio_rs": [653_587.46],
        }
    ).write_parquet(iso_dir / "arm2_alpha_grid.parquet")
    _write_item_metrics(
        iso_dir / "item_decomposition_alpha066.parquet",
        margem_realizada_rs=4_765_904.95,
        ruptura_rs=264_998.27,
    )
    return results


@pytest.mark.usefixtures("cenario")
def test_n_months_no_periodo() -> None:
    assert mr.n_months_no_periodo() == pytest.approx(270 / 30)


@pytest.mark.usefixtures("cenario")
def test_mensalizar_divide_pelo_numero_de_meses() -> None:
    assert mr.mensalizar(90.0) == pytest.approx(10.0)  # 90 / (270/30) = 10


def test_ruptura_evitada_e_positiva_quando_o_pior_tem_mais_ruptura() -> None:
    pior = mr.ArmFinancials("pior", 0.6, 1.0, 100.0, 0.0, 1000.0, 0.0, 1)
    melhor = mr.ArmFinancials("melhor", 0.95, 1.0, 200.0, 0.0, 200.0, 0.0, 1)
    assert mr.ruptura_evitada(pior, melhor) == pytest.approx(800.0)


def test_capital_liberado_e_negativo_quando_recalibrar_aumenta_capital() -> None:
    # Sprint 19: recalibrar EXIGE mais capital (não libera) -- ver docstring
    # do módulo e a aba "Andar 1" do relatório.
    desatualizado = mr.ArmFinancials("desatualizado", 0.6, 17.0, 175_000.0, 0.0, 0.0, 0.0, 1)
    recalibrado = mr.ArmFinancials("recalibrado", 0.95, 7.0, 672_000.0, 0.0, 0.0, 0.0, 1)
    aumento = mr.capital_liberado(recalibrado, desatualizado)
    assert aumento > 0  # aumento de capital, não liberação
    assert mr.capital_liberado(desatualizado, recalibrado) == pytest.approx(-aumento)


def test_perda_evitada_e_zero_quando_ambos_tem_perda_zero() -> None:
    a = mr.ArmFinancials("a", 0.9, 1.0, 100.0, 0.0, 0.0, 0.0, 1)
    b = mr.ArmFinancials("b", 0.9, 1.0, 100.0, 0.0, 0.0, 0.0, 1)
    assert mr.perda_evitada(a, b) == 0.0


def test_sensibilidade_margem_escala_linearmente() -> None:
    faixa = mr.sensibilidade_margem(1000.0)
    assert faixa[0.7] == pytest.approx(700.0)
    assert faixa[1.0] == pytest.approx(1000.0)
    assert faixa[1.3] == pytest.approx(1300.0)


@pytest.mark.usefixtures("cenario")
def test_load_motor_financials_soma_item_metrics_e_le_portfolio_do_grid() -> None:
    motor = mr.load_motor_financials()
    assert motor.nivel_servico == pytest.approx(0.9483648)
    assert motor.decision_metric == pytest.approx(7.29191612)
    assert motor.margem_realizada_rs == pytest.approx(4_765_904.95)
    assert motor.ruptura_rs == pytest.approx(264_998.27)
    assert motor.n_items == 2


@pytest.mark.usefixtures("cenario")
def test_load_recalibrado_financials() -> None:
    recalibrado = mr.load_recalibrado_financials()
    assert recalibrado.nivel_servico == pytest.approx(0.9485436)
    assert recalibrado.ruptura_rs == pytest.approx(265_066.16)


@pytest.mark.usefixtures("cenario")
def test_load_stale_financials_para_as_duas_magnitudes() -> None:
    lt3 = mr.load_stale_financials("calibrado_lt3_rp7")
    lt5 = mr.load_stale_financials("calibrado_lt5_rp7")
    assert lt3.nivel_servico == pytest.approx(0.5999)
    assert lt5.nivel_servico == pytest.approx(0.6635)
    assert lt3.ruptura_rs != lt5.ruptura_rs


@pytest.mark.usefixtures("cenario")
def test_ruptura_evitada_andar1_bate_com_numeros_medidos_na_sprint_18() -> None:
    stale = mr.load_stale_financials("calibrado_lt3_rp7")
    recalibrado = mr.load_recalibrado_financials()
    evitada = mr.ruptura_evitada(stale, recalibrado)
    assert evitada == pytest.approx(2_019_048.41 - 265_066.16, abs=1.0)


# --------------------------------------------------------------------------
# aba "Resumo visual" (Sprint 23)
# --------------------------------------------------------------------------


def _visual_sheet(tmp_path: Path):
    path = tmp_path / "visual.xlsx"
    workbook = xlsxwriter.Workbook(str(path))
    mr._write_visual_sheet(workbook, mr._build_report_data(), mr._make_formats(workbook))
    workbook.close()
    return load_workbook(path)["Resumo visual"]


def _find(ws, text: str) -> int:
    for row in ws.iter_rows():
        if row[0].value is not None and text in str(row[0].value):
            return row[0].row
    raise AssertionError(f"texto {text!r} não encontrado na aba")


@pytest.mark.usefixtures("cenario")
def test_visual_tres_blocos_com_numeros_calculados_dos_resultados(tmp_path: Path) -> None:
    ws = _visual_sheet(tmp_path)
    stale = mr.load_stale_financials(mr.STALE_LABEL_HEADLINE)
    recal = mr.load_recalibrado_financials()
    motor = mr.load_motor_financials()

    servico = str(ws.cell(row=_find(ws, "NÍVEL DE SERVIÇO") + 1, column=1).value)
    assert "60%" in servico
    assert "95%" in servico

    capital = ws.cell(row=_find(ws, "CAPITAL PARADO") + 1, column=1).value
    assert capital == pytest.approx(mr.variacao_relativa_capital(stale, recal))
    assert capital == pytest.approx(2.84, abs=0.01)  # +284%

    margem = ws.cell(row=_find(ws, "MARGEM POR REAL") + 1, column=1).value
    assert margem == pytest.approx(motor.decision_metric / recal.decision_metric - 1)
    assert margem == pytest.approx(0.0288, abs=0.0005)  # +2,9%


@pytest.mark.usefixtures("cenario")
def test_visual_dados_dos_graficos_sao_os_numeros_dos_dois_graficos(tmp_path: Path) -> None:
    ws = _visual_sheet(tmp_path)
    top = _find(ws, "Dados dos gráficos")
    rows = [[c.value for c in ws[r]][:3] for r in range(top + 2, top + 5)]
    assert [r[0] for r in rows] == ["ERP desatualizado", "ERP recalibrado", "Motor (braço 2)"]
    assert [round(r[1], 4) for r in rows] == [0.5999, 0.9485, 0.9484]  # serviço, fração
    assert [round(r[2], 1) for r in rows] == [26.0, 100.0, 97.2]  # capital, 100 = recalibrado


@pytest.mark.usefixtures("cenario")
def test_visual_tem_dois_graficos_pequenos_empilhados(tmp_path: Path) -> None:
    ws = _visual_sheet(tmp_path)
    assert len(ws._charts) == 2
    anchors = [c.anchor._from for c in ws._charts]
    assert anchors[0].col == anchors[1].col  # empilhados na mesma coluna
    assert anchors[0].row < anchors[1].row


@pytest.mark.usefixtures("cenario")
def test_visual_declara_celula_origem_e_ilustrativo(tmp_path: Path) -> None:
    ws = _visual_sheet(tmp_path)
    texto = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "lead time 7" in texto and "revisão 14" in texto
    assert "results/baseline_desatualizado" in texto
    assert "results/iso_servico" in texto
    assert "ilustrativo" in texto  # R$ das outras abas


def test_workbook_abre_na_aba_visual_e_mantem_as_demais_atras(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, cenario: Path
) -> None:
    """As abas que leem o canônico real são trocadas por abas vazias de mesmo nome: o que se
    testa aqui é a ORDEM e a aba ativa, não o conteúdo delas."""
    for name, sheet_name in (
        ("_write_sangrando_sheet", "Itens sangrando margem"),
        ("_write_premissas_sheet", "Premissas"),
    ):
        monkeypatch.setattr(
            mr, name, lambda wb, *_a, _n=sheet_name, **_k: wb.add_worksheet(_n), raising=True
        )
    path = mr.write_monthly_report_workbook(tmp_path / "r.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames[:3] == [
        "Resumo visual",
        "Resumo executivo",
        "Andar 1 - Correção de cadastro",
    ]
    assert wb.active.title == "Resumo visual"
    assert cenario.exists()

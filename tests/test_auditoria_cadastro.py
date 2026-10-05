"""Testes de `motor.reporting.auditoria_cadastro` (Sprint 22).

A auditoria é a primeira entrega que acha dinheiro sem modelo (documento de
negócio, seção 5). Demonstrada sobre dado FICTÍCIO: os achados esperados são
EXATAMENTE os problemas plantados em `tests/fixtures/make_fixtures.py` --
conjuntos comparados por igualdade, nada a mais e nada a menos.
"""

from __future__ import annotations

import dataclasses
from pathlib import Path

import pytest
from openpyxl import load_workbook

from motor.io.loaders_cliente import ClientLoadResult, load_client_tables
from motor.reporting.auditoria_cadastro import build_cadastre_audit, write_audit_workbook
from tests.cliente_helpers import FIXTURE_DIR, STORE_CNPJ, STORE_ID, client_params


@pytest.fixture(scope="module")
def result() -> ClientLoadResult:
    return load_client_tables(
        FIXTURE_DIR, params=client_params(), store_id=STORE_ID, store_cnpj=STORE_CNPJ
    )


@pytest.fixture(scope="module")
def audit(result: ClientLoadResult):
    return build_cadastre_audit(result, params=client_params())


def test_margem_negativa(audit) -> None:
    assert set(audit.negative_margin["item_id"]) == {"1003"}
    linha = audit.negative_margin.row(0, named=True)
    assert linha["price_ref"] == pytest.approx(4.50)
    assert linha["cost"] == pytest.approx(5.20)
    assert linha["margin_pct"] == pytest.approx((4.50 - 5.20) / 4.50)


def test_ean_duplicado_distingue_o_que_foi_unido_do_que_nao_foi(audit) -> None:
    rows = {(r["ean"], tuple(r["item_ids"]), r["merged"]) for r in audit.duplicate_ean.to_dicts()}
    assert rows == {
        ("7891000000081", ("1008", "1009"), True),  # mesma unidade: unido (DC4)
        ("7891000000111", ("1011", "1012"), False),  # unidades diferentes: só apontado
    }


def test_codigo_repetido_no_cadastro(audit) -> None:
    assert set(audit.repeated_codes["item_id"]) == {"1013"}


def test_toda_uniao_de_codigos_aparece_na_auditoria(audit) -> None:
    assert audit.code_merges.select("survivor", "merged", "ean").to_dicts() == [
        {"survivor": "1008", "merged": "1009", "ean": "7891000000081"}
    ]


def test_custo_de_cadastro_diferente_da_ultima_nota(audit) -> None:
    # 1001: 17,00 vs 18,00 (+5,9%) e 1006: 1,20 vs 1,25 (+4,2%) ficam DENTRO da
    # tolerância de 10%; só o 1002 (8,00 vs 9,80, +22,5%) estoura
    assert set(audit.cost_divergence["item_id"]) == {"1002"}
    linha = audit.cost_divergence.row(0, named=True)
    assert linha["cadastro_cost"] == pytest.approx(8.00)
    assert linha["last_nfe_cost"] == pytest.approx(9.80)
    assert linha["divergence"] == pytest.approx(9.80 / 8.00 - 1)


def test_parado_com_saldo(audit) -> None:
    assert audit.idle_stock_evaluated is True
    assert set(audit.idle_stock["item_id"]) == {"1014"}
    linha = audit.idle_stock.row(0, named=True)
    assert linha["on_hand"] == 40.0
    assert linha["days_since_last_sale"] == 100  # 2024-01-20 -> 2024-04-29


def test_sem_stock_a_checagem_diz_nao_avaliada_em_vez_de_zero_achados(
    result: ClientLoadResult,
) -> None:
    sem_stock = dataclasses.replace(result, stock=result.stock.clear())
    audit = build_cadastre_audit(sem_stock, params=client_params())
    assert audit.idle_stock_evaluated is False
    assert audit.idle_stock.height == 0


def test_fardo_que_a_nota_nao_informa(audit) -> None:
    rows = {r["item_id"]: r["pack_source"] for r in audit.pack_unknown.to_dicts()}
    assert rows == {"1005": "cadastro", "1006": "fornecedor"}


def test_workbook_tem_uma_aba_por_checagem(audit, tmp_path: Path) -> None:
    path = write_audit_workbook(audit, tmp_path / "auditoria.xlsx", illustrative=True)
    wb = load_workbook(path)
    assert wb.sheetnames == [
        "Resumo",
        "Margem negativa",
        "EAN e códigos duplicados",
        "Custo divergente",
        "Parado com saldo",
        "Fardo",
        "Validade não informada",
    ]
    resumo = {r[0].value: r[1].value for r in wb["Resumo"].iter_rows() if r[0].value}
    assert resumo["Margem negativa (itens)"] == 1
    assert resumo["Custo de cadastro ≠ última nota (itens)"] == 1
    assert resumo["Parado com saldo (itens)"] == 1
    margem = [r[0].value for r in wb["Margem negativa"].iter_rows(min_row=2) if r[0].value]
    assert "1003" in margem


def test_workbook_declara_que_e_demonstracao_sobre_dado_ficticio(audit, tmp_path: Path) -> None:
    path = write_audit_workbook(audit, tmp_path / "a.xlsx", illustrative=True)
    textos = [c.value for row in load_workbook(path)["Resumo"].iter_rows() for c in row if c.value]
    assert any("exemplo" in str(t).lower() and "fictício" in str(t).lower() for t in textos)


def test_workbook_sem_stock_escreve_nao_avaliado(result: ClientLoadResult, tmp_path: Path) -> None:
    audit = build_cadastre_audit(
        dataclasses.replace(result, stock=result.stock.clear()), params=client_params()
    )
    path = write_audit_workbook(audit, tmp_path / "a.xlsx", illustrative=False)
    ws = load_workbook(path)["Parado com saldo"]
    textos = [c.value for row in ws.iter_rows() for c in row if c.value]
    assert any("não avaliado" in str(t).lower() for t in textos)


def test_perecivel_sem_validade_informada(audit) -> None:
    # 1007 informa 5 dias; 1011 e 1012 são perecíveis sem validade_dias
    assert set(audit.perishable_no_shelf_life["item_id"]) == {"1011", "1012"}

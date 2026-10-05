"""Testes do modo `canonical` de `generate_purchase_list` (Sprint 22, fase 2).

Prova a promessa da Sprint 1 no ponto mais exposto: a lista de compra roda
sobre as tabelas do cliente (`motor.io.loaders_cliente`) lendo DELAS fardo,
pedido mínimo, lead time, revisão, estoque e pedidos em aberto -- e não os
valores assumidos do Favorita nem o estoque que o simulador inventa.

Os comportamentos exigidos (cada um com teste):

1. `on_hand` vem da tabela `stock` na véspera da decisão; sem `stock` a
   lista só sai como DEMONSTRAÇÃO ("estoque simulado, não usar para pedido
   real" no topo de cada aba e na aba Notas).
2. `in_transit` vem de `pedidos_em_aberto.csv`; sem o arquivo, é 0 -- declarado
   na aba Notas -- e fornecedor com lead_time >= revisão vai inteiro para
   exceções ("pedido anterior pode estar em trânsito").
3. Sem orçamento em `client_loader.weekly_budget_rs`, a restrição de caixa fica
   desligada, declarado na aba Notas.

O modo Favorita (`assumed_cell`) tem regressão própria em
`tests/test_purchase_list_regression.py`.
"""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

import polars as pl
import pytest
from openpyxl import load_workbook

from motor.config import Params, load_params
from motor.experiments.iso_service import CanonicalTables
from motor.io.loaders_cliente import ClientLoadResult, client_params_for, load_client_tables
from motor.policy.guardrails import GuardrailReason
from motor.reporting.purchase_list import (
    CanonicalInputs,
    PurchaseListRow,
    SupplierPurchaseList,
    SupplierTerms,
    generate_purchase_list,
    write_supplier_workbook,
)
from tests.cliente_helpers import (
    FIXTURE_DIR,
    S1,
    S4,
    STORE_CNPJ,
    STORE_ID,
    client_params,
)

PARAMS_PATH = Path(__file__).resolve().parent.parent / "config" / "params.yaml"
THURSDAY = date(2024, 5, 2)  # S1 pede segunda/quinta; saldo de 2024-04-29 é o mais recente
TUESDAY = date(2024, 4, 30)  # S4 pede terça
IN_SERIES = date(2024, 4, 25)  # dentro da série de vendas -- o simulador do modo demo precisa
KEEP = object()

IN_TRANSIT_WARNING = "pedido anterior pode estar em trânsito"


@pytest.fixture(scope="module")
def result() -> ClientLoadResult:
    return load_client_tables(
        FIXTURE_DIR, params=client_params(), store_id=STORE_ID, store_cnpj=STORE_CNPJ
    )


@pytest.fixture(scope="module")
def params(result: ClientLoadResult) -> Params:
    return client_params_for(load_params(PARAMS_PATH), result)


def _run(
    result: ClientLoadResult,
    params: Params,
    tmp_path: Path,
    as_of: date,
    *,
    sales: pl.DataFrame | None = None,
    stock: pl.DataFrame | None = None,
    open_orders: object = KEEP,
    purchases: object = KEEP,
) -> dict[str, SupplierPurchaseList]:
    tables = CanonicalTables(
        sales=(result.sales if sales is None else sales).lazy(),
        items=result.items,
        suppliers=result.suppliers,
        stock=result.stock if stock is None else stock,
        subset_cache_path=tmp_path / "subset.json",
    )
    inputs = CanonicalInputs(
        purchases=result.purchases if purchases is KEEP else purchases,  # type: ignore[arg-type]
        open_orders=result.open_orders if open_orders is KEEP else open_orders,  # type: ignore[arg-type]
    )
    lists, _report = generate_purchase_list(
        as_of,
        base_params=params,
        tables=tables,
        supplier_terms=SupplierTerms.CANONICAL,
        canonical_inputs=inputs,
    )
    return lists


def _rows(lst: SupplierPurchaseList) -> dict[str, PurchaseListRow]:
    return {r.item_id: r for r in (*lst.main, *lst.exceptions)}


# --------------------------------------------------------------------------
# termos de fornecedor lidos do canônico
# --------------------------------------------------------------------------


def test_fardo_pedido_minimo_e_prazos_vem_das_tabelas_do_cliente(
    result: ClientLoadResult, params: Params, tmp_path: Path
) -> None:
    lst = _run(result, params, tmp_path, THURSDAY)[S1]
    assert set(_rows(lst)) == {
        "1001", "1002", "1003", "1004", "1005", "1006", "1007", "1008", "1011", "1012",
        "1014", "1017",
    }  # fmt: skip
    assert lst.min_order_value == pytest.approx(
        500.0
    )  # do fornecedor, não do dicionário por categoria
    assert (lst.lead_time_days, lst.review_period_days) == (3, 4)  # não o 7/14 do Favorita
    rows = _rows(lst)
    for item_id, fardo in (("1001", 6.0), ("1002", 12.0), ("1004", 12.0), ("1005", 24.0)):
        row = rows[item_id]
        assert row.quantity_units % fardo == 0.0, item_id  # sempre múltiplo do fardo do cliente
        assert row.quantity_packs == pytest.approx(row.quantity_units / fardo)


def test_fornecedor_com_varias_categorias_nao_levanta_erro(
    result: ClientLoadResult, params: Params, tmp_path: Path
) -> None:
    lst = _run(result, params, tmp_path, THURSDAY)[S1]  # o S1 tem 5 categorias
    assert {r.category for r in _rows(lst).values()} >= {
        "GROCERY",
        "BEVERAGES",
        "DAIRY",
        "CLEANING",
    }


def test_descricao_do_produto_na_linha(
    result: ClientLoadResult, params: Params, tmp_path: Path
) -> None:
    rows = _rows(_run(result, params, tmp_path, THURSDAY)[S1])
    assert rows["1001"].description == "ARROZ TIPO 1 5KG"
    assert rows["1007"].description == "CARNE MOIDA KG"


# --------------------------------------------------------------------------
# 1. on_hand real da tabela stock; sem stock, demonstração
# --------------------------------------------------------------------------


def test_on_hand_vem_da_tabela_stock_nao_do_simulador(
    result: ClientLoadResult, params: Params, tmp_path: Path
) -> None:
    lst = _run(result, params, tmp_path, THURSDAY)[S1]
    on_hand = {i: r.on_hand for i, r in _rows(lst).items()}
    assert on_hand == {
        "1001": 30.0, "1002": 50.0, "1003": 80.0, "1004": 100.0, "1005": 40.0, "1006": 20.0,
        "1007": 12.5, "1008": 25.0, "1011": 3.0, "1012": 2.0, "1014": 40.0, "1017": 0.0,
    }  # fmt: skip
    assert lst.provenance is not None
    assert lst.provenance.simulated_stock is False
    assert lst.provenance.stock_date == date(2024, 4, 29)


def test_item_sem_linha_no_stock_vai_para_excecoes_com_saldo_zero(
    result: ClientLoadResult, params: Params, tmp_path: Path
) -> None:
    stock = result.stock.filter(pl.col("item_id") != "1001")
    row = _rows(_run(result, params, tmp_path, THURSDAY, stock=stock)[S1])["1001"]
    assert row.on_hand == 0.0
    assert row.is_exception
    assert any("saldo de estoque não informado" in a for a in row.adjustments)


def test_sem_stock_a_lista_e_demonstracao_com_estoque_simulado(
    result: ClientLoadResult, params: Params, tmp_path: Path
) -> None:
    lst = _run(result, params, tmp_path, IN_SERIES, stock=result.stock.clear())[S1]
    assert lst.provenance is not None
    assert lst.provenance.simulated_stock is True
    assert lst.provenance.stock_date is None
    assert lst.lead_time_days == 3  # o simulador usa o prazo do cliente, não o 7/14


# --------------------------------------------------------------------------
# 2. in_transit dos pedidos em aberto; sem arquivo, 0 + lead >= revisão -> exceções
# --------------------------------------------------------------------------


def test_in_transit_vem_dos_pedidos_em_aberto(
    result: ClientLoadResult, params: Params, tmp_path: Path
) -> None:
    lst = _run(result, params, tmp_path, THURSDAY)[S1]
    in_transit = {i: r.in_transit for i, r in _rows(lst).items()}
    assert in_transit["1001"] == 60.0
    assert in_transit["1002"] == 120.0
    assert in_transit["1007"] == 25.5
    assert in_transit["1003"] == 0.0  # arquivo existe e não lista o item: sabe-se que é zero
    assert lst.provenance is not None
    assert lst.provenance.in_transit_source == "pedidos_em_aberto"


def test_sem_arquivo_de_pedidos_in_transit_e_zero_e_e_declarado(
    result: ClientLoadResult, params: Params, tmp_path: Path
) -> None:
    lst = _run(result, params, tmp_path, THURSDAY, open_orders=None)[S1]
    assert all(r.in_transit == 0.0 for r in _rows(lst).values())
    assert lst.provenance is not None
    assert lst.provenance.in_transit_source == "zero_sem_arquivo"
    # S1: lead 3 < revisão 4 -- sem o arquivo, ainda assim nada vai para exceção por isso
    assert not any(IN_TRANSIT_WARNING in a for r in _rows(lst).values() for a in r.adjustments)


def test_lead_time_maior_ou_igual_a_revisao_sem_arquivo_manda_o_fornecedor_para_excecoes(
    result: ClientLoadResult, params: Params, tmp_path: Path
) -> None:
    lst = _run(result, params, tmp_path, TUESDAY, open_orders=None)[S4]
    assert (lst.lead_time_days, lst.review_period_days) == (7, 7)
    assert lst.main == ()  # TODOS os itens do fornecedor vão para exceções
    assert set(_rows(lst)) == {"1015", "1016"}
    for row in lst.exceptions:
        assert any(IN_TRANSIT_WARNING in a for a in row.adjustments), row.item_id


def test_com_o_arquivo_de_pedidos_o_mesmo_fornecedor_nao_e_suspeito(
    result: ClientLoadResult, params: Params, tmp_path: Path
) -> None:
    """O S4 não tem nenhum pedido em aberto, mas o ARQUIVO existe: sabe-se que nada
    está a caminho. Sem o arquivo seria dúvida; com ele, é informação."""
    lst = _run(result, params, tmp_path, TUESDAY)[S4]
    assert not any(IN_TRANSIT_WARNING in a for r in _rows(lst).values() for a in r.adjustments)
    assert all(r.in_transit == 0.0 for r in _rows(lst).values())


# --------------------------------------------------------------------------
# compras reais, item novo, caixa
# --------------------------------------------------------------------------


def test_limite_de_variacao_usa_as_compras_reais_das_notas(
    result: ClientLoadResult, params: Params, tmp_path: Path
) -> None:
    tiny = pl.DataFrame(
        {"item_id": ["1001"], "date": [date(2024, 4, 20)], "units": [1.0]},
        schema={"item_id": pl.String, "date": pl.Date, "units": pl.Float64},
    )
    row = _rows(_run(result, params, tmp_path, THURSDAY, purchases=tiny)[S1])["1001"]
    assert row.is_exception
    assert any("média das compras" in a for a in row.adjustments)


def test_sem_historico_de_compras_o_limite_de_variacao_fica_desligado(
    result: ClientLoadResult, params: Params, tmp_path: Path
) -> None:
    lst = _run(result, params, tmp_path, THURSDAY, purchases=None)[S1]
    assert not any("média das compras" in a for r in _rows(lst).values() for a in r.adjustments)
    assert lst.provenance is not None
    assert lst.provenance.purchases_source == "ausente"


def test_item_com_pouco_historico_e_item_sem_historico_vao_pela_regra_de_item_novo(
    result: ClientLoadResult, params: Params, tmp_path: Path
) -> None:
    sales = result.sales.filter(
        ~((pl.col("item_id") == "1015") & (pl.col("date") < date(2024, 4, 10)))
        & (pl.col("item_id") != "1016")
    )
    lst = _run(result, params, tmp_path, TUESDAY, sales=sales)[S4]
    rows = _rows(lst)
    assert set(rows) == {"1015", "1016"}  # nenhum item é cortado por "pouco histórico"
    for item_id in ("1015", "1016"):
        assert rows[item_id].is_exception
        assert any("item novo" in a for a in rows[item_id].adjustments), item_id


def test_restricao_de_caixa_desligada_sem_orcamento_e_ligada_com_ele(
    result: ClientLoadResult, params: Params, tmp_path: Path
) -> None:
    off = _run(result, params, tmp_path, THURSDAY)[S1]
    assert off.provenance is not None
    assert off.provenance.cash_constraint_active is False
    assert not any("caixa" in a for r in _rows(off).values() for a in r.adjustments)

    budgeted = params.model_copy(
        update={"client_loader": params.client_loader.model_copy(update={"weekly_budget_rs": 1.0})}
    )
    on = _run(result, budgeted, tmp_path, THURSDAY)[S1]
    assert on.provenance is not None
    assert on.provenance.cash_constraint_active is True
    assert any("caixa" in a for r in _rows(on).values() for a in r.adjustments)
    assert GuardrailReason.RESTRICAO_DE_CAIXA.value == "restricao_de_caixa"


# --------------------------------------------------------------------------
# entradas inválidas
# --------------------------------------------------------------------------


def test_params_sem_validade_da_categoria_do_cliente_levanta_erro_claro(
    result: ClientLoadResult, tmp_path: Path
) -> None:
    with pytest.raises(ValueError, match="client_params_for"):
        _run(result, load_params(PARAMS_PATH), tmp_path, THURSDAY)


def test_modo_favorita_recusa_insumos_canonicos(
    result: ClientLoadResult, params: Params, tmp_path: Path
) -> None:
    tables = CanonicalTables(
        sales=result.sales.lazy(),
        items=result.items,
        suppliers=result.suppliers,
        stock=result.stock,
        subset_cache_path=tmp_path / "s.json",
    )
    with pytest.raises(ValueError, match="canonical_inputs"):
        generate_purchase_list(
            THURSDAY,
            base_params=params,
            tables=tables,
            supplier_terms=SupplierTerms.ASSUMED_CELL,
            canonical_inputs=CanonicalInputs(purchases=None, open_orders=None),
        )


# --------------------------------------------------------------------------
# xlsx: banner de demonstração, Notas, coluna de descrição
# --------------------------------------------------------------------------


def _cells(path: Path, sheet: str) -> list[str]:
    return [str(c.value) for row in load_workbook(path)[sheet].iter_rows() for c in row if c.value]


def test_xlsx_do_modo_demo_traz_o_aviso_no_topo_de_cada_aba_e_nas_notas(
    result: ClientLoadResult, params: Params, tmp_path: Path
) -> None:
    lst = _run(result, params, tmp_path, IN_SERIES, stock=result.stock.clear())[S1]
    path = write_supplier_workbook(lst, as_of=IN_SERIES, out_dir=tmp_path / "x")
    wb = load_workbook(path)
    for sheet in ("Lista principal", "Exceções", "Notas"):
        first = str(wb[sheet].cell(row=1, column=1).value)
        assert "ESTOQUE SIMULADO" in first, sheet
        assert "não usar para pedido real" in first.lower(), sheet
    assert any("estoque simulado" in c.lower() for c in _cells(path, "Notas")[1:])


def test_xlsx_com_estoque_real_nao_tem_aviso_e_tem_descricao_e_notas_declaradas(
    result: ClientLoadResult, params: Params, tmp_path: Path
) -> None:
    lst = _run(result, params, tmp_path, THURSDAY, open_orders=None, purchases=None)[S1]
    path = write_supplier_workbook(lst, as_of=THURSDAY, out_dir=tmp_path / "x")
    wb = load_workbook(path)
    header = [str(c.value) for c in wb["Lista principal"][1]]
    assert header[:2] == ["código", "descrição"]
    assert "ESTOQUE SIMULADO" not in " ".join(_cells(path, "Lista principal"))
    notes = " ".join(_cells(path, "Notas")).lower()
    assert "pedidos em aberto" in notes  # in_transit = 0 declarado
    assert "restrição de caixa desligada" in notes
    assert "limite de variação desligado" in notes
    assert "combinação não medida" in notes  # alpha=0,66 foi calibrado na célula 7/14
    assert "frete" not in header  # (sanidade: nenhuma coluna inventada)


def test_xlsx_com_dados_completos_nao_declara_o_que_nao_falta(
    result: ClientLoadResult, params: Params, tmp_path: Path
) -> None:
    budgeted = params.model_copy(
        update={"client_loader": params.client_loader.model_copy(update={"weekly_budget_rs": 1e9})}
    )
    lst = _run(result, budgeted, tmp_path, THURSDAY)[S1]
    notes = " ".join(
        _cells(write_supplier_workbook(lst, as_of=THURSDAY, out_dir=tmp_path / "x"), "Notas")
    ).lower()
    assert "restrição de caixa desligada" not in notes
    assert "limite de variação desligado" not in notes


# --------------------------------------------------------------------------
# item parado e item novo descontando a posição
# --------------------------------------------------------------------------


def test_item_parado_com_historico_recebe_sugestao_zero_e_vai_para_excecoes(
    result: ClientLoadResult, params: Params, tmp_path: Path
) -> None:
    """1014: vendeu até 2024-01-20, tem saldo 40. Tem histórico (não é item novo) mas está
    parado há mais de `idle_days_threshold` dias: não se pede mais -- sugestão 0."""
    row = _rows(_run(result, params, tmp_path, THURSDAY)[S1])["1014"]
    assert row.quantity_units == 0.0
    assert row.value_rs == 0.0
    assert row.is_exception
    assert any("item parado: 103 dias sem venda, saldo 40" in a for a in row.adjustments), (
        row.adjustments
    )
    assert not any("item novo" in a for a in row.adjustments)  # parado vence item novo


def test_item_parado_nao_entra_na_mediana_da_categoria_do_item_novo(
    result: ClientLoadResult, params: Params, tmp_path: Path
) -> None:
    # 1017 também está parado (última venda 2024-02-01) e é do S1 junto do 1014: os dois saem
    # com 0 e nenhum deles vira "referência" da categoria para um item novo
    rows = _rows(_run(result, params, tmp_path, THURSDAY)[S1])
    assert rows["1017"].quantity_units == 0.0
    assert any("item parado" in a for a in rows["1017"].adjustments)


def test_item_novo_desconta_a_posicao(
    result: ClientLoadResult, params: Params, tmp_path: Path
) -> None:
    """pedido = max(0, mediana - (on_hand + in_transit)): o item novo não pede o que já tem."""
    sales = result.sales.filter(
        ~((pl.col("item_id") == "1015") & (pl.col("date") < date(2024, 4, 10)))
    )
    lst = _run(result, params, tmp_path, TUESDAY, sales=sales)[S4]
    row = _rows(lst)["1015"]  # saldo 5
    detail = next(a for a in row.adjustments if "item novo" in a)
    median = float(re.search(r"mediana da categoria \((\d+\.\d+) unidades", detail).group(1))  # type: ignore[union-attr]
    assert "menos a posição de 5.00" in detail
    assert row.on_hand == 5.0
    # pedido = mediana - posição (20,61), arredondado para cima no fardo -- nunca a mediana cheia
    assert median - 5.0 <= row.quantity_units < median

    # posição que já cobre a mediana: não pede nada
    stock = result.stock.with_columns(
        pl.when(pl.col("item_id") == "1015")
        .then(500.0)
        .otherwise(pl.col("on_hand"))
        .alias("on_hand")
    )
    covered = _rows(_run(result, params, tmp_path, TUESDAY, sales=sales, stock=stock)[S4])["1015"]
    assert covered.quantity_units == 0.0

"""Testes de `motor.reporting.purchase_week` (Sprint 22): a lista semanal num arquivo só.

Um xlsx da semana com uma aba por fornecedor, em ordem de dia de pedido
(depois `supplier_id`), mais a aba inicial "Resumo da semana". O xlsx por
fornecedor (`write_supplier_workbook`) continua existindo; este é adicional.

Cada fornecedor aparece uma vez, na SUA PRIMEIRA data de pedido da semana --
fornecedor com dois dias de pedido (segunda e quinta) tem a segunda decisão
apontada no Resumo, não duplicada em outra aba.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from motor.config import load_params
from motor.experiments.iso_service import CanonicalTables
from motor.io.loaders_cliente import ClientLoadResult, client_params_for, load_client_tables
from motor.policy.supplier import OrderOutcome
from motor.reporting.purchase_list import (
    DEMO_BANNER,
    CanonicalInputs,
    ListProvenance,
    PurchaseListRow,
    SupplierPurchaseList,
    SupplierTerms,
)
from motor.reporting.purchase_week import (
    generate_weekly_lists,
    safe_sheet_name,
    write_weekly_workbook,
)
from tests.cliente_helpers import FIXTURE_DIR, S1, S3, S4, STORE_CNPJ, STORE_ID, client_params

PARAMS_PATH = Path(__file__).resolve().parent.parent / "config" / "params.yaml"
MON, TUE, WED, THU = date(2024, 5, 6), date(2024, 5, 7), date(2024, 5, 8), date(2024, 5, 9)
assert MON.isoweekday() == 1


def _row(
    item_id: str, *, value: float, exception: bool = False, reason: str = ""
) -> PurchaseListRow:
    return PurchaseListRow(
        item_id=item_id,
        category="CAT",
        unit_of_sale="unidade",
        quantity_units=10.0,
        quantity_packs=10.0,
        on_hand=1.0,
        in_transit=0.0,
        coverage_days_after_order=7.0,
        value_rs=value,
        margin_pct=0.3,
        adjustments=(reason,) if reason else (),
        is_exception=exception,
        description=f"PRODUTO {item_id}",
    )


def _list(supplier_id: str, *, demo: bool = False, exceptions: bool = True) -> SupplierPurchaseList:
    provenance = ListProvenance(
        supplier_terms=SupplierTerms.CANONICAL,
        simulated_stock=demo,
        stock_date=None if demo else date(2024, 4, 29),
        in_transit_source="simulado" if demo else "pedidos_em_aberto",
        purchases_source="simulado" if demo else "notas",
        cash_constraint_active=False,
    )
    return SupplierPurchaseList(
        supplier_id=supplier_id,
        category="CAT",
        outcome=OrderOutcome.PEDIDO_EMITIDO,
        min_order_value=0.0,
        main=(_row(f"{supplier_id}-A", value=100.0), _row(f"{supplier_id}-B", value=50.0)),
        exceptions=(
            (_row(f"{supplier_id}-X", value=25.0, exception=True, reason="motivo de teste"),)
            if exceptions
            else ()
        ),
        lead_time_days=3,
        review_period_days=7,
        provenance=provenance,
    )


def _names(path: Path) -> list[str]:
    return load_workbook(path).sheetnames


def _cells(path: Path, sheet: str) -> list[str]:
    return [str(c.value) for row in load_workbook(path)[sheet].iter_rows() for c in row if c.value]


# --------------------------------------------------------------------------
# ordem e conteúdo
# --------------------------------------------------------------------------


def test_abas_em_ordem_de_dia_de_pedido_depois_supplier_id(tmp_path: Path) -> None:
    by_date = {
        WED: {"A": _list("A")},
        MON: {"C": _list("C"), "B": _list("B")},  # mesma data: desempate por supplier_id
        TUE: {},
    }
    path = write_weekly_workbook(by_date, tmp_path / "semana.xlsx", week_start=MON)
    assert _names(path) == ["Resumo da semana", "B", "C", "A"]


def test_resumo_lista_na_mesma_ordem_com_data_dia_da_semana_e_totais(tmp_path: Path) -> None:
    by_date = {MON: {"B": _list("B")}, WED: {"A": _list("A", exceptions=False)}}
    path = write_weekly_workbook(by_date, tmp_path / "s.xlsx", week_start=MON)
    ws = load_workbook(path)["Resumo da semana"]
    rows = [[c.value for c in row] for row in ws.iter_rows() if row[0].value in ("B", "A")]
    assert [r[0] for r in rows] == ["B", "A"]
    assert str(rows[0][1]).startswith("2024-05-06") and "segunda" in str(rows[0][1])
    assert "quarta" in str(rows[1][1])
    # colunas: fornecedor, dia do pedido, resultado, itens (principal), exceções, total R$
    assert (rows[0][3], rows[0][4], rows[0][5]) == (2, 1, 175.0)
    assert (rows[1][3], rows[1][4], rows[1][5]) == (2, 0, 150.0)


def test_aba_do_fornecedor_traz_principal_e_excecoes_com_motivo(tmp_path: Path) -> None:
    path = write_weekly_workbook({MON: {"B": _list("B")}}, tmp_path / "s.xlsx", week_start=MON)
    cells = _cells(path, "B")
    assert "Lista principal" in cells and "Exceções" in cells
    assert "B-A" in cells and "B-X" in cells
    assert "motivo de teste" in cells
    assert "PRODUTO B-A" in cells  # coluna de descrição


def test_segunda_data_de_pedido_do_mesmo_fornecedor_e_apontada_no_resumo(tmp_path: Path) -> None:
    by_date = {MON: {"B": _list("B")}}
    path = write_weekly_workbook(
        by_date, tmp_path / "s.xlsx", week_start=MON, other_order_dates={"B": (THU,)}
    )
    texts = " ".join(_cells(path, "Resumo da semana"))
    assert "2024-05-09" in texts  # a quinta: próxima decisão do B, fora desta aba


def test_aviso_de_demonstracao_no_topo_de_todas_as_abas(tmp_path: Path) -> None:
    by_date = {MON: {"B": _list("B", demo=True)}, TUE: {"C": _list("C", demo=True)}}
    path = write_weekly_workbook(by_date, tmp_path / "s.xlsx", week_start=MON)
    wb = load_workbook(path)
    for name in wb.sheetnames:
        first = str(wb[name].cell(row=1, column=1).value)
        assert DEMO_BANNER in first, name


def test_sem_demonstracao_nao_ha_aviso(tmp_path: Path) -> None:
    path = write_weekly_workbook({MON: {"B": _list("B")}}, tmp_path / "s.xlsx", week_start=MON)
    assert not any("ESTOQUE SIMULADO" in c for n in _names(path) for c in _cells(path, n))


def test_resumo_declara_que_a_restricao_de_caixa_e_por_data_de_decisao(tmp_path: Path) -> None:
    path = write_weekly_workbook({MON: {"B": _list("B")}}, tmp_path / "s.xlsx", week_start=MON)
    texts = " ".join(_cells(path, "Resumo da semana")).lower()
    assert "por data de decisão" in texts
    assert "restrição de caixa" in texts


def test_semana_sem_nenhum_pedido_gera_so_o_resumo(tmp_path: Path) -> None:
    path = write_weekly_workbook({}, tmp_path / "s.xlsx", week_start=MON)
    assert _names(path) == ["Resumo da semana"]
    assert any("nenhum fornecedor" in c.lower() for c in _cells(path, "Resumo da semana"))


@pytest.mark.parametrize(
    ("raw", "esperado"),
    [
        ("12345678000101", "12345678000101"),
        ("FORN/A:B*C?[x]", "FORN_A_B_C__x_"),
        ("X" * 40, "X" * 31),
    ],
)
def test_nome_de_aba_valido_no_excel(raw: str, esperado: str) -> None:
    assert safe_sheet_name(raw, taken=set()) == esperado


def test_nome_de_aba_repetido_ganha_sufixo_unico() -> None:
    taken = {"X" * 31}
    out = safe_sheet_name("X" * 40, taken=taken)
    assert out not in taken
    assert len(out) <= 31


# --------------------------------------------------------------------------
# da tabela do cliente ao arquivo semanal
# --------------------------------------------------------------------------


@pytest.fixture(scope="module")
def result() -> ClientLoadResult:
    return load_client_tables(
        FIXTURE_DIR, params=client_params(), store_id=STORE_ID, store_cnpj=STORE_CNPJ
    )


def test_semana_do_cliente_roda_o_pipeline_uma_vez_por_fornecedor_e_ordena_as_abas(
    result: ClientLoadResult, tmp_path: Path
) -> None:
    params = client_params_for(load_params(PARAMS_PATH), result)
    tables = CanonicalTables(
        sales=result.sales.lazy(),
        items=result.items,
        suppliers=result.suppliers,
        stock=result.stock,
        subset_cache_path=tmp_path / "unused.json",
    )
    by_date, other_dates = generate_weekly_lists(
        MON,
        base_params=params,
        tables=tables,
        supplier_terms=SupplierTerms.CANONICAL,
        canonical_inputs=CanonicalInputs(
            purchases=result.purchases, open_orders=result.open_orders
        ),
    )
    # S1 pede seg/qui, S4 pede ter, S3 pede qua: cada um aparece na sua PRIMEIRA data
    assert {d: sorted(v) for d, v in by_date.items() if v} == {MON: [S1], TUE: [S4], WED: [S3]}
    assert other_dates == {S1: (THU,)}
    path = write_weekly_workbook(
        by_date, tmp_path / "semana.xlsx", week_start=MON, other_order_dates=other_dates
    )
    assert _names(path) == ["Resumo da semana", S1, S4, S3]

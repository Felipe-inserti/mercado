"""Regressão do modo Favorita (`assumed_cell`) de `generate_purchase_list` (Sprint 22).

A Sprint 22 abre um modo `canonical` na lista de compra (dado de cliente). A
condição para isso é que o modo ORIGINAL continue produzindo a MESMA saída,
linha a linha e célula a célula. Este teste compara a saída atual com um
golden (`tests/fixtures/golden_purchase_list_assumed_cell.json`) capturado do
código ANTES da mudança, sobre um cenário sintético com categorias do
Favorita (determinístico, sem depender de dado real).

Regenerar o golden só se a mudança de saída for DECIDIDA (nunca para "fazer o
teste passar"): `python -m tests.test_purchase_list_regression`.
"""

from __future__ import annotations

import json
import random
from dataclasses import asdict
from datetime import date, timedelta
from pathlib import Path

import polars as pl
from openpyxl import load_workbook

from motor.config import load_params
from motor.experiments.iso_service import CanonicalTables
from motor.reporting.purchase_list import generate_purchase_list, write_supplier_workbook

PARAMS_PATH = Path(__file__).resolve().parent.parent / "config" / "params.yaml"
GOLDEN_PATH = (
    Path(__file__).resolve().parent / "fixtures" / "golden_purchase_list_assumed_cell.json"
)
START, END = date(2024, 1, 1), date(2024, 4, 29)
AS_OF = date(2024, 4, 18)  # quinta: dia de pedido dos fornecedores abaixo
CATEGORIES = ("GROCERY I", "BEVERAGES", "CLEANING")


def build_scenario(tmp_dir: Path):
    """Tabelas canônicas sintéticas com cara de Favorita: loja '44', um
    fornecedor por categoria (`SUP-<CATEGORY>`), `pack_multiple` 1,0 e sem
    estoque -- como o canônico do Favorita."""
    rng = random.Random(7)
    items, sales, suppliers = [], [], []
    n_days = (END - START).days + 1
    for c_idx, category in enumerate(CATEGORIES):
        supplier_id = f"SUP-{category.replace(' ', '_')}"
        suppliers.append(
            {
                "supplier_id": supplier_id, "lead_time_days": 3, "order_days": [1, 4],
                "review_period_days": 7, "min_order_value": 0.0, "min_order_units": 0.0,
            }
        )  # fmt: skip
        for k in range(4):
            item_id = f"{100 * (c_idx + 1) + k}"
            items.append(
                {
                    "item_id": item_id, "ean": None, "description": None, "category": category,
                    "item_class": None, "supplier_id": supplier_id, "unit_of_sale": "unidade",
                    "pack_multiple": 1.0, "is_perishable": False, "is_anchor": None,
                    "cost": 6.0 + k, "price_ref": 8.0 + k, "economics_origin": "arbitrado",
                }
            )  # fmt: skip
            for d in range(n_days):
                sales.append(
                    {
                        "store_id": "44", "item_id": item_id, "date": START + timedelta(days=d),
                        "units_sold": float(rng.randint(5, 25) + c_idx), "units_returned": 0.0,
                        "on_promo": False, "price": None, "is_operating_day": True,
                        "is_anomaly": False, "anomaly_reason": None,
                    }
                )  # fmt: skip
    stock = pl.DataFrame(
        schema={
            "store_id": pl.String, "item_id": pl.String, "date": pl.Date,
            "on_hand": pl.Float64, "in_transit": pl.Float64,
        }
    )  # fmt: skip
    base = load_params(PARAMS_PATH)
    sel = base.subset_selection.model_copy(
        update={
            "store_id": "44", "n_items_min": 5, "n_items_max": 30, "min_weeks_of_history": 6,
            "density_window_start": START, "density_window_end": date(2024, 2, 25),
            "family_min_items": 1, "family_cap": 30,
        }
    )  # fmt: skip
    params = base.model_copy(
        update={
            "subset_selection": sel,
            "canonical": base.canonical.model_copy(
                update={"window_start": START, "window_end": END, "anomalies": []}
            ),
            "simulation": base.simulation.model_copy(
                update={
                    "warmup_days": 14, "start_date": date(2024, 2, 26),
                    "evaluation_start_date": date(2024, 3, 11), "end_date": END,
                }
            ),
        }
    )  # fmt: skip
    tables = CanonicalTables(
        sales=pl.DataFrame(sales).lazy(),
        items=pl.DataFrame(items),
        suppliers=pl.DataFrame(suppliers),
        stock=stock,
        subset_cache_path=tmp_dir / "subset.json",
    )
    return params, tables


def _original_fields(row: object) -> dict[str, object]:
    """Campos de `PurchaseListRow` anteriores à Sprint 22. `description` (campo novo)
    não entra no golden -- é exigido `None` aqui: o canônico do Favorita não tem
    descrição, e a coluna só aparece quando há descrição."""
    fields = asdict(row)  # type: ignore[call-overload]
    assert fields.pop("description") is None
    return fields


def snapshot(tmp_dir: Path) -> dict[str, object]:
    """Saída completa: linhas por fornecedor, contagem de guardas e o conteúdo
    de TODAS as células de cada xlsx."""
    params, tables = build_scenario(tmp_dir)
    results, report = generate_purchase_list(AS_OF, base_params=params, tables=tables)
    out: dict[str, object] = {
        "report_counts": {
            k.value: v for k, v in sorted(report.counts.items(), key=lambda kv: kv[0].value)
        }
    }
    suppliers: dict[str, object] = {}
    for supplier_id, lst in sorted(results.items()):
        path = write_supplier_workbook(lst, as_of=AS_OF, out_dir=tmp_dir / "xlsx")
        wb = load_workbook(path)
        suppliers[supplier_id] = {
            "outcome": str(lst.outcome),
            "category": lst.category,
            "min_order_value": lst.min_order_value,
            "main": [_original_fields(r) for r in lst.main],
            "exceptions": [_original_fields(r) for r in lst.exceptions],
            "xlsx": {
                ws.title: [[c.value for c in row] for row in ws.iter_rows()] for ws in wb.worksheets
            },
        }
    out["suppliers"] = suppliers
    return json.loads(json.dumps(out, default=str))


def test_modo_favorita_mantem_a_saida_identica(tmp_path: Path) -> None:
    golden = json.loads(GOLDEN_PATH.read_text(encoding="utf-8"))
    assert snapshot(tmp_path) == golden


if __name__ == "__main__":
    import tempfile

    with tempfile.TemporaryDirectory() as tmp:
        data = snapshot(Path(tmp))
    GOLDEN_PATH.write_text(
        json.dumps(data, indent=1, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    print(f"golden gravado em {GOLDEN_PATH}")
